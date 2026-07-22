"""Bounded streaming readers for CSV and XLSX completed-sales files."""

import csv
import re
from collections.abc import Iterator, Sequence
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from marktwert.application.imports import RawSalesRow, SalesFileError, SalesFileReader

MAXIMUM_SOURCE_FILE_BYTES = 100 * 1024 * 1024
MAXIMUM_XLSX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
MAXIMUM_XLSX_ARCHIVE_ENTRIES = 20_000
MAXIMUM_XLSX_COMPRESSION_RATIO = 200
CSV_SNIFF_BYTES = 16_384
SUPPORTED_FILE_EXTENSIONS = frozenset({".csv", ".xlsx"})


class TabularSalesFileReaderFactory:
    """Validate a path and select its bounded streaming reader."""

    def create(self, path: Path) -> SalesFileReader:
        """Return a reader for a supported, safe local file."""
        resolved_path = path.expanduser().resolve()
        if not resolved_path.is_file():
            message = f"import file does not exist: {resolved_path}"
            raise SalesFileError(message)
        file_size = resolved_path.stat().st_size
        if file_size > MAXIMUM_SOURCE_FILE_BYTES:
            message = "import file exceeds the 100 MB safety limit"
            raise SalesFileError(message)

        suffix = resolved_path.suffix.casefold()
        if suffix not in SUPPORTED_FILE_EXTENSIONS:
            message = "supported import formats are CSV and XLSX"
            raise SalesFileError(message)
        if suffix == ".xlsx":
            _validate_xlsx_archive(resolved_path)
            return XlsxSalesFileReader()
        return CsvSalesFileReader()


class CsvSalesFileReader:
    """Stream UTF-8 CSV rows with delimiter and header normalization."""

    def read_rows(self, path: Path) -> Iterator[RawSalesRow]:
        """Yield normalized CSV rows."""
        try:
            with path.open(
                "r",
                encoding="utf-8-sig",
                errors="strict",
                newline="",
            ) as file:
                sample = file.read(CSV_SNIFF_BYTES)
                file.seek(0)
                dialect = _detect_csv_dialect(sample)
                reader = csv.reader(file, dialect)
                try:
                    raw_headers = next(reader)
                except StopIteration as error:
                    message = "CSV file is empty"
                    raise SalesFileError(message) from error
                headers = _normalize_headers(raw_headers)
                for row_number, row in enumerate(reader, start=2):
                    yield RawSalesRow(
                        row_number=row_number,
                        values=_row_values(headers, row),
                    )
        except UnicodeDecodeError as error:
            message = "CSV file must use UTF-8 encoding"
            raise SalesFileError(message) from error


class XlsxSalesFileReader:
    """Stream values from the first worksheet of a safe XLSX workbook."""

    def read_rows(self, path: Path) -> Iterator[RawSalesRow]:
        """Yield normalized XLSX rows without loading the workbook into memory."""
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration as error:
                message = "XLSX worksheet is empty"
                raise SalesFileError(message) from error
            headers = _normalize_headers(raw_headers)
            for row_number, row in enumerate(rows, start=2):
                yield RawSalesRow(
                    row_number=row_number,
                    values=_row_values(headers, row),
                )
        finally:
            workbook.close()


def _detect_csv_dialect(sample: str) -> type[csv.Dialect] | csv.Dialect:
    if not sample.strip():
        message = "CSV file is empty"
        raise SalesFileError(message)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        return csv.excel


def _normalize_headers(headers: Sequence[object | None]) -> tuple[str, ...]:
    normalized = tuple(_normalize_header(value) for value in headers)
    if not normalized or any(not header for header in normalized):
        message = "every import column must have a header"
        raise SalesFileError(message)
    if len(set(normalized)) != len(normalized):
        message = "import column headers must be unique"
        raise SalesFileError(message)
    return normalized


def _normalize_header(value: object | None) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"_+", "_", re.sub(r"[^\w]+", "_", text.casefold())).strip("_")


def _row_values(
    headers: tuple[str, ...],
    row: Sequence[object | None],
) -> dict[str, object | None]:
    extra_values = row[len(headers) :]
    if any(value not in (None, "") for value in extra_values):
        message = "row contains more values than the header"
        raise SalesFileError(message)
    return {
        header: row[index] if index < len(row) else None
        for index, header in enumerate(headers)
    }


def _validate_xlsx_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAXIMUM_XLSX_ARCHIVE_ENTRIES:
                message = "XLSX archive contains too many entries"
                raise SalesFileError(message)
            uncompressed_size = sum(entry.file_size for entry in entries)
            compressed_size = sum(max(entry.compress_size, 1) for entry in entries)
            if uncompressed_size > MAXIMUM_XLSX_UNCOMPRESSED_BYTES:
                message = "XLSX expanded content exceeds the 250 MB safety limit"
                raise SalesFileError(message)
            if (
                uncompressed_size
                > compressed_size * MAXIMUM_XLSX_COMPRESSION_RATIO
            ):
                message = "XLSX compression ratio exceeds the safety limit"
                raise SalesFileError(message)
            if any(entry.flag_bits & 0x1 for entry in entries):
                message = "encrypted XLSX files are not supported"
                raise SalesFileError(message)
    except BadZipFile as error:
        message = "XLSX file is not a valid Office archive"
        raise SalesFileError(message) from error
